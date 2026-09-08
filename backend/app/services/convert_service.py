"""
Layanan konversi dokumen untuk UltraPDF.

Semua konversi berbasis LibreOffice memakai pool profil yang dipakai ulang
(lihat LibreOfficeProfilePool) supaya konversi kedua dan seterusnya tidak
membayar ongkos pembuatan profil baru seperti pada PDFService.
"""

import asyncio
import csv
import io
import logging
import os
import re
import shutil
import tempfile
import threading
import uuid
import zipfile
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Callable, Optional

from app.services.pdf_service import PDFService

logger = logging.getLogger(__name__)

PROCESS_TIMEOUT = int(os.getenv("PROCESS_TIMEOUT", "300"))
LIBREOFFICE_POOL_SIZE = int(os.getenv("LIBREOFFICE_POOL_SIZE", "2"))
LIBREOFFICE_PROFILE_ROOT = os.getenv(
    "LIBREOFFICE_PROFILE_ROOT",
    os.path.join(tempfile.gettempdir(), "ultrapdf_lo_profiles"),
)

# Ambang deteksi lapisan teks PDF, dipakai saat konversi PDF ke Word.
# Di bawah PDF_MIN_CHARS_PER_PAGE halaman dianggap hasil pindaian murni; di atas
# PDF_HIDDEN_TEXT_RATIO teksnya dianggap lapisan OCR tak terlihat; DOCX dianggap
# gagal bila teksnya kurang dari PDF_DOCX_TEXT_RATIO kali teks PDF aslinya.
PDF_MIN_CHARS_PER_PAGE = int(os.getenv("PDF_MIN_CHARS_PER_PAGE", "20"))
PDF_HIDDEN_TEXT_RATIO = float(os.getenv("PDF_HIDDEN_TEXT_RATIO", "0.6"))
PDF_DOCX_TEXT_RATIO = float(os.getenv("PDF_DOCX_TEXT_RATIO", "0.5"))

# Pengaturan OCR untuk PDF hasil pindaian. OCR jauh lebih lambat daripada
# konversi biasa (sekitar 2-3 detik per halaman), jadi batas waktunya dihitung
# per halaman, bukan memakai PROCESS_TIMEOUT yang datar.
OCR_DPI = int(os.getenv("OCR_DPI", "300"))
OCR_LANGUAGE = os.getenv("OCR_LANGUAGE", "ind+eng")
OCR_TIMEOUT_PER_PAGE = float(os.getenv("OCR_TIMEOUT_PER_PAGE", "15"))
OCR_MAX_PAGES = int(os.getenv("OCR_MAX_PAGES", "300"))

# Deteksi gambar halaman (logo kop surat, tanda tangan, stempel) pada PDF hasil
# pindaian. Halaman dirender pada GRAPHIC_DPI untuk mencari areanya, lalu
# potongannya dirender ulang pada GRAPHIC_RENDER_DPI supaya tetap tajam di Word.
GRAPHIC_DPI = int(os.getenv("OCR_GRAPHIC_DPI", "150"))
GRAPHIC_RENDER_DPI = int(os.getenv("OCR_GRAPHIC_RENDER_DPI", "220"))
GRAPHIC_INK_THRESHOLD = int(os.getenv("OCR_GRAPHIC_INK_THRESHOLD", "200"))
GRAPHIC_MIN_AREA_PT = float(os.getenv("OCR_GRAPHIC_MIN_AREA_PT", "40"))
GRAPHIC_MAX_REGIONS = int(os.getenv("OCR_GRAPHIC_MAX_REGIONS", "12"))

# Pembersihan halaman hasil foto: bidang kertas dicari pada citra yang
# diperkecil ke PHOTO_EDGE_WORK_SIZE piksel, kemiringan diukur pada
# PHOTO_SKEW_WORK_SIZE piksel, dan cahaya baru diratakan kalau selisih terang
# gelap latarnya melewati PHOTO_SHADOW_SPREAD.
PHOTO_EDGE_WORK_SIZE = int(os.getenv("PHOTO_EDGE_WORK_SIZE", "900"))
PHOTO_SKEW_WORK_SIZE = int(os.getenv("PHOTO_SKEW_WORK_SIZE", "1200"))
PHOTO_QUAD_MIN_AREA = float(os.getenv("PHOTO_QUAD_MIN_AREA", "0.3"))
PHOTO_QUAD_MAX_AREA = float(os.getenv("PHOTO_QUAD_MAX_AREA", "0.97"))
PHOTO_SHADOW_SPREAD = int(os.getenv("PHOTO_SHADOW_SPREAD", "25"))
DESKEW_MIN_ANGLE = float(os.getenv("DESKEW_MIN_ANGLE", "0.3"))
DESKEW_MAX_ANGLE = float(os.getenv("DESKEW_MAX_ANGLE", "12"))

# Tinggi kotak hasil OCR dikali angka ini untuk memperkirakan ukuran huruf.
OCR_FONT_HEIGHT_RATIO = float(os.getenv("OCR_FONT_HEIGHT_RATIO", "0.78"))

# Ekstensi yang bisa dibuka LibreOffice dan diekspor ke PDF
OFFICE_INPUT_EXTENSIONS = {
    ".doc", ".docx", ".odt", ".rtf", ".txt",
    ".xls", ".xlsx", ".ods", ".csv",
    ".ppt", ".pptx", ".odp",
    ".html", ".htm",
}

# Filter impor LibreOffice untuk format yang ambigu
LIBREOFFICE_INPUT_FILTERS = {
    ".csv": "Text - txt - csv (StarCalc)",
    ".html": "HTML (StarWriter)",
    ".htm": "HTML (StarWriter)",
}

# Filter ekspor PDF untuk tiap keluarga dokumen (dipakai saat butuh opsi khusus)
PDF_EXPORT_FILTERS = {
    "writer": "writer_pdf_Export",
    "calc": "calc_pdf_Export",
    "impress": "impress_pdf_Export",
    "draw": "draw_pdf_Export",
}

# Ekstensi gambar yang bisa dihasilkan dari PDF. PyMuPDF hanya bisa menulis
# JPG dan PNG sendiri; WebP dan TIFF ditulis lewat Pillow.
IMAGE_TARGET_EXTENSIONS = {
    "jpg": "jpg",
    "jpeg": "jpg",
    "png": "png",
    "webp": "webp",
    "tif": "tiff",
    "tiff": "tiff",
}

PILLOW_IMAGE_FORMATS = {"webp": "WEBP", "tiff": "TIFF"}


def resolve_image_extension(image_format: str) -> str:
    """Normalkan nama format gambar menjadi ekstensi berkas yang dipakai."""
    extension = IMAGE_TARGET_EXTENSIONS.get((image_format or "").lower())
    if not extension:
        allowed = ", ".join(sorted(set(IMAGE_TARGET_EXTENSIONS)))
        raise ValueError(f"Format gambar harus salah satu dari: {allowed}")
    return extension


SPREADSHEET_EXTENSIONS = {".xls", ".xlsx", ".ods", ".csv"}
PRESENTATION_EXTENSIONS = {".ppt", ".pptx", ".odp"}


class LibreOfficeProfilePool:
    """
    Pool profil pengguna LibreOffice.

    PDFService membuat direktori profil baru tiap request lalu menghapusnya.
    Itu aman tapi lambat karena LibreOffice harus menginisialisasi profil dari nol
    setiap konversi. Pool ini menyimpan sejumlah profil yang dipakai bergantian,
    sekaligus membatasi jumlah proses LibreOffice yang berjalan bersamaan.
    """

    def __init__(
        self,
        size: int = LIBREOFFICE_POOL_SIZE,
        root: str = LIBREOFFICE_PROFILE_ROOT,
    ):
        self.size = max(1, size)
        self.root = root
        self._queue: Optional[asyncio.Queue] = None
        self._lock = asyncio.Lock()

    async def _ensure_queue(self) -> asyncio.Queue:
        # Queue dibuat lazily karena butuh event loop yang sudah berjalan
        if self._queue is not None:
            return self._queue

        async with self._lock:
            if self._queue is None:
                os.makedirs(self.root, exist_ok=True)
                queue: asyncio.Queue = asyncio.Queue()
                for index in range(self.size):
                    profile_dir = os.path.join(self.root, f"profile_{index}")
                    os.makedirs(profile_dir, exist_ok=True)
                    queue.put_nowait(profile_dir)
                self._queue = queue
                logger.info(
                    f"LibreOffice profile pool siap ({self.size} profil di {self.root})"
                )

        return self._queue

    @asynccontextmanager
    async def acquire(self):
        queue = await self._ensure_queue()
        profile_dir = await queue.get()
        try:
            os.makedirs(profile_dir, exist_ok=True)
            yield profile_dir
        finally:
            queue.put_nowait(profile_dir)

    async def reset(self, profile_dir: str):
        """Buang profil yang korup (dipanggil setelah konversi gagal)."""
        try:
            await asyncio.to_thread(shutil.rmtree, profile_dir, True)
            os.makedirs(profile_dir, exist_ok=True)
            logger.warning(f"Profil LibreOffice direset: {profile_dir}")
        except Exception as e:
            logger.error(f"Gagal mereset profil LibreOffice {profile_dir}: {e}")


libreoffice_pool = LibreOfficeProfilePool()


def parse_page_ranges(spec: Optional[str], total_pages: int) -> list[int]:
    """
    Ubah spesifikasi halaman ("1-3,5,8-") menjadi daftar indeks 0-based.
    Spec kosong berarti semua halaman.
    """
    if not spec or not spec.strip():
        return list(range(total_pages))

    pages: list[int] = []
    for chunk in spec.split(","):
        chunk = chunk.strip()
        if not chunk:
            continue

        if "-" in chunk:
            start_raw, _, end_raw = chunk.partition("-")
            try:
                start = int(start_raw) if start_raw.strip() else 1
                end = int(end_raw) if end_raw.strip() else total_pages
            except ValueError:
                raise ValueError(f"Rentang halaman tidak valid: {chunk}")
        else:
            try:
                start = end = int(chunk)
            except ValueError:
                raise ValueError(f"Nomor halaman tidak valid: {chunk}")

        if start < 1 or end < start:
            raise ValueError(f"Rentang halaman tidak valid: {chunk}")

        for page in range(start, min(end, total_pages) + 1):
            if page - 1 not in pages:
                pages.append(page - 1)

    if not pages:
        raise ValueError("Rentang halaman tidak menghasilkan halaman apa pun")

    return pages


def create_zip(file_paths: list[str], output_path: str) -> str:
    """Bungkus beberapa hasil konversi menjadi satu ZIP."""
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    used_names: dict[str, int] = {}
    with zipfile.ZipFile(output_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in file_paths:
            arcname = os.path.basename(path)
            # Cegah tabrakan nama di dalam arsip
            if arcname in used_names:
                used_names[arcname] += 1
                stem, ext = os.path.splitext(arcname)
                arcname = f"{stem}_{used_names[arcname]}{ext}"
            else:
                used_names[arcname] = 0
            archive.write(path, arcname)

    return output_path


class PageCanvas:
    """
    Citra satu halaman PDF yang siap dipakai untuk OCR dan pemotongan gambar.

    PDF hasil foto ponsel tidak pernah rapi: kertasnya miring, sisinya menjorok
    karena perspektif, dan cahayanya tidak rata. Kanvas ini memegang versi
    halaman yang sudah diluruskan beserta ukuran kertas hasil pelurusannya,
    sehingga OCR, deteksi gambar, dan potongan logo maupun tanda tangan semuanya
    memakai satu citra yang sama. Tanpa citra (OpenCV tidak terpasang) semua
    permintaan jatuh kembali ke render langsung dari halaman PDF-nya.
    """

    def __init__(self, page, image=None, width_pt: float = 0.0, height_pt: float = 0.0):
        self.page = page
        self.image = image
        self.width_pt = width_pt or page.rect.width
        self.height_pt = height_pt or page.rect.height

    @property
    def scale(self) -> float:
        """Jumlah piksel citra untuk tiap titik halaman."""
        if self.image is None:
            return OCR_DPI / 72.0
        return self.image.shape[1] / max(self.width_pt, 1.0)

    def to_png(self) -> bytes:
        """Seluruh halaman sebagai PNG, untuk disodorkan ke mesin OCR."""
        if self.image is None:
            return self.page.get_pixmap(dpi=OCR_DPI).tobytes("png")

        import cv2

        ok, buffer = cv2.imencode(".png", self.image)
        if not ok:
            return self.page.get_pixmap(dpi=OCR_DPI).tobytes("png")
        return buffer.tobytes()

    def gray(self):
        """Citra halaman dalam abu-abu, atau None bila citranya tidak ada."""
        if self.image is None:
            return None

        import cv2

        if self.image.ndim == 2:
            return self.image
        return cv2.cvtColor(self.image, cv2.COLOR_BGR2GRAY)

    def crop_png(self, bbox) -> Optional[bytes]:
        """Potong satu kotak halaman (satuan titik) menjadi PNG."""
        if self.image is None:
            fitz = ConvertService._import_fitz()

            rect = fitz.Rect(*bbox) + (-2.0, -2.0, 2.0, 2.0)
            rect = rect & self.page.rect
            if rect.is_empty:
                return None
            return self.page.get_pixmap(dpi=GRAPHIC_RENDER_DPI, clip=rect).tobytes("png")

        import cv2

        scale = self.scale
        height, width = self.image.shape[:2]
        left = max(int(round((bbox[0] - 2.0) * scale)), 0)
        top = max(int(round((bbox[1] - 2.0) * scale)), 0)
        right = min(int(round((bbox[2] + 2.0) * scale)), width)
        bottom = min(int(round((bbox[3] + 2.0) * scale)), height)
        if right <= left or bottom <= top:
            return None

        ok, buffer = cv2.imencode(".png", self.image[top:bottom, left:right])
        return buffer.tobytes() if ok else None


class ConvertService:
    # ------------------------------------------------------------------
    # Office (Word/Excel/PowerPoint/ODF/RTF/TXT/HTML) -> PDF
    # ------------------------------------------------------------------
    @staticmethod
    async def office_to_pdf(
        input_path: str,
        output_dir: str,
        pdf_variant: Optional[str] = None,
        fit_to_page: bool = False,
    ) -> Optional[str]:
        """
        Konversi dokumen Office/ODF/teks apa pun yang dikenal LibreOffice ke PDF.

        pdf_variant "pdfa" mengekspor sebagai PDF/A-2b.
        fit_to_page hanya berlaku untuk spreadsheet (semua kolom masuk satu halaman).
        """
        if not os.path.exists(input_path):
            logger.error(f"Input file not found: {input_path}")
            return None

        os.makedirs(output_dir, exist_ok=True)
        ext = Path(input_path).suffix.lower()

        convert_target = "pdf"
        export_options: list[str] = []

        if pdf_variant == "pdfa":
            # SelectPdfVersion 2 = PDF/A-2b pada LibreOffice 7+
            export_options.append('"SelectPdfVersion":{"type":"long","value":"2"}')

        if fit_to_page and ext in SPREADSHEET_EXTENSIONS:
            export_options.append('"SinglePageSheets":{"type":"boolean","value":"true"}')

        if export_options:
            if ext in SPREADSHEET_EXTENSIONS:
                family = "calc"
            elif ext in PRESENTATION_EXTENSIONS:
                family = "impress"
            else:
                family = "writer"
            joined = ",".join(export_options)
            convert_target = f"pdf:{PDF_EXPORT_FILTERS[family]}:{{{joined}}}"

        command_tail = ["--convert-to", convert_target, "--outdir", output_dir, input_path]

        infilter = LIBREOFFICE_INPUT_FILTERS.get(ext)
        if infilter:
            command_tail.insert(0, f"--infilter={infilter}")

        return await ConvertService._run_libreoffice(
            input_path, output_dir, command_tail, ".pdf", "Office to PDF"
        )

    # ------------------------------------------------------------------
    # PDF -> Office
    # ------------------------------------------------------------------
    @staticmethod
    def _analyze_pdf_text(input_path: str) -> dict:
        """
        Periksa lapisan teks PDF supaya mode konversi ke DOCX bisa dipilih.

        PDF hasil pemindaian yang sudah di-OCR menyimpan teksnya sebagai teks
        tak terlihat (render mode 3) di atas gambar halaman. pdf2docx secara
        bawaan membuang teks semacam itu, sehingga DOCX-nya hanya berisi gambar
        dan tidak bisa diedit. Hasil analisis ini yang menentukan penanganannya.
        """
        fitz = ConvertService._import_fitz()

        document = fitz.open(input_path)
        try:
            page_count = document.page_count or 1
            total_chars = 0
            visible_chars = 0
            hidden_chars = 0

            for page in document:
                total_chars += len(page.get_text("text").strip())
                try:
                    spans = page.get_texttrace()
                except Exception:
                    # get_texttrace bisa gagal pada font rusak; abaikan halamannya
                    continue

                for span in spans:
                    char_count = len(span.get("chars") or ())
                    if span.get("type") == 3:
                        hidden_chars += char_count
                    else:
                        visible_chars += char_count
        finally:
            document.close()

        traced = visible_chars + hidden_chars
        return {
            "pages": page_count,
            "chars": total_chars,
            "chars_per_page": total_chars / page_count,
            "hidden_ratio": (hidden_chars / traced) if traced else 0.0,
        }

    @staticmethod
    def _docx_text_length(path: str) -> int:
        """Hitung panjang teks yang benar-benar bisa diedit di dalam DOCX."""
        try:
            from docx import Document
        except ImportError:
            # Tanpa python-docx hasilnya tidak bisa diverifikasi; anggap lolos
            return -1

        def paragraphs_length(paragraphs) -> int:
            return sum(len(paragraph.text.strip()) for paragraph in paragraphs)

        try:
            document = Document(path)
        except Exception:
            return 0

        total = paragraphs_length(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    total += paragraphs_length(cell.paragraphs)
        return total

    @staticmethod
    def _run_pdf2docx(input_path: str, output_path: str, ocr_mode: int) -> None:
        """
        Jalankan pdf2docx dengan mode lapisan teks tertentu.

        ocr_mode 0 memakai teks yang terlihat (PDF digital biasa), ocr_mode 2
        memakai teks tak terlihat hasil OCR dan mengabaikan gambar halaman.
        """
        from pdf2docx import Converter

        converter = Converter(input_path)
        try:
            converter.convert(output_path, start=0, end=None, ocr=ocr_mode)
        finally:
            converter.close()

    @staticmethod
    def _build_docx_from_pdf_text(input_path: str, output_path: str) -> None:
        """
        Bangun DOCX langsung dari span teks PDF memakai python-docx.

        Cadangan terakhir ketika pdf2docx gagal menghasilkan teks: layout kolom
        dan tabel tidak dipertahankan, tapi seluruh isinya dijamin berupa teks
        yang bisa diedit, lengkap dengan ukuran huruf, tebal, dan miringnya.
        """
        fitz = ConvertService._import_fitz()
        from docx import Document
        from docx.shared import Pt

        document = fitz.open(input_path)
        docx_document = Document()
        try:
            for page_index, page in enumerate(document):
                if page_index:
                    docx_document.add_page_break()

                for block in page.get_text("dict").get("blocks", []):
                    if block.get("type") != 0:  # 0 = blok teks
                        continue

                    lines = block.get("lines") or []
                    if not any(
                        span.get("text", "").strip()
                        for line in lines
                        for span in line.get("spans") or ()
                    ):
                        continue

                    paragraph = docx_document.add_paragraph()
                    for line_index, line in enumerate(lines):
                        if line_index:
                            paragraph.add_run(" ")
                        for span in line.get("spans") or ():
                            text = span.get("text", "")
                            if not text:
                                continue
                            run = paragraph.add_run(text)
                            size = span.get("size")
                            if size:
                                run.font.size = Pt(round(float(size), 1))
                            flags = span.get("flags", 0)
                            run.bold = bool(flags & 2 ** 4)
                            run.italic = bool(flags & 2 ** 1)

            docx_document.save(output_path)
        finally:
            document.close()

    _ocr_reader = None
    _ocr_reader_lock = threading.Lock()

    @staticmethod
    def _load_ocr_reader():
        """
        Ambil mesin OCR yang tersedia, atau None kalau tidak ada satu pun.

        Dikembalikan sebagai fungsi yang menerima PNG halaman dan mengembalikan
        daftar baris teks sesuai urutan baca. RapidOCR dipakai lebih dulu karena
        cukup dipasang lewat pip (memakai onnxruntime yang sudah jadi dependensi),
        sedangkan Tesseract butuh binary sistem.

        Hasilnya di-cache karena memuat model ONNX makan waktu beberapa detik;
        tanpa cache ongkos itu dibayar ulang tiap konversi.
        """
        if ConvertService._ocr_reader is not None:
            return ConvertService._ocr_reader or None

        with ConvertService._ocr_reader_lock:
            if ConvertService._ocr_reader is None:
                ConvertService._ocr_reader = ConvertService._build_ocr_reader() or False
        return ConvertService._ocr_reader or None

    @staticmethod
    def _build_ocr_reader():
        """
        Siapkan pembaca OCR dari mesin pertama yang terpasang.

        Pembaca menerima PNG halaman dan mengembalikan daftar dict berisi
        ``text`` dan ``bbox`` (x0, y0, x1, y1 dalam piksel gambar). Kotaknya ikut
        dikembalikan karena fitur "PDF bisa dicari" perlu menempatkan teks tak
        terlihat tepat di atas tulisan aslinya, bukan sekadar tahu isinya.
        """
        try:
            from rapidocr_onnxruntime import RapidOCR
        except ImportError:
            pass
        else:
            engine = RapidOCR()

            def read_with_rapidocr(image_bytes: bytes) -> list:
                # use_cls=False: klasifikator sudut kadang membalik baris yang
                # sebenarnya sudah tegak sehingga isinya rusak. Halaman di sini
                # selalu dirender dari PDF (rotasi halaman sudah diterapkan
                # PyMuPDF), jadi orientasinya dijamin benar dan pemeriksaan itu
                # hanya menambah risiko sekaligus memperlambat.
                # return_word_box=True menambah kotak tiap karakter. Itu dipakai
                # untuk mengembalikan spasi yang kadang tidak ikut dikenali pada
                # tulisan kapital seperti nama instansi di kop surat.
                result, _ = engine(image_bytes, use_cls=False, return_word_box=True)
                if not result:
                    return []

                # Tiap entri: [4 titik kotak, teks, skor, kotak karakter, karakter, skor]
                items = []
                for entry in result:
                    box, text = entry[0], str(entry[1])
                    if len(entry) >= 5:
                        respaced = ConvertService._respace_ocr_text(entry[4], entry[3])
                        if respaced:
                            text = respaced
                    if not text.strip():
                        continue
                    xs = [point[0] for point in box]
                    ys = [point[1] for point in box]
                    items.append(
                        {
                            "text": text,
                            "bbox": (min(xs), min(ys), max(xs), max(ys)),
                        }
                    )
                return items

            return read_with_rapidocr

        try:
            import io

            import pytesseract
            from PIL import Image
        except ImportError:
            return None

        def read_with_tesseract(image_bytes: bytes) -> list:
            with Image.open(io.BytesIO(image_bytes)) as image:
                data = pytesseract.image_to_data(
                    image, lang=OCR_LANGUAGE, output_type=pytesseract.Output.DICT
                )

            items = []
            for index, text in enumerate(data["text"]):
                if not text.strip():
                    continue
                left = data["left"][index]
                top = data["top"][index]
                items.append(
                    {
                        "text": text,
                        "bbox": (
                            left,
                            top,
                            left + data["width"][index],
                            top + data["height"][index],
                        ),
                    }
                )
            return items

        return read_with_tesseract

    @staticmethod
    def _respace_ocr_text(chars, boxes) -> Optional[str]:
        """
        Kembalikan spasi yang hilang dari hasil pengenalan satu baris.

        Model pengenal kerap menyatukan kata, apalagi pada foto yang agak buram:
        nama instansi di kop surat terbaca seperti "DINASTEKNOLOGIINFORMASI".
        Kotak tiap karakter masih memuat jaraknya, jadi celah yang jauh lebih
        lebar daripada jarak antarhuruf biasa dikembalikan menjadi spasi. Spasi
        yang sudah dikenali tidak diutak-atik.
        """
        if not chars or not boxes or len(chars) != len(boxes) or len(chars) < 3:
            return None
        if any(len(character) != 1 for character in chars):
            return None

        spans = []
        for box in boxes:
            try:
                xs = [float(point[0]) for point in box]
            except (TypeError, ValueError, IndexError):
                return None
            spans.append((min(xs), max(xs)))

        # Kotak spasi tidak ikut dihitung: lebarnya bukan lebar huruf
        widths = sorted(
            max(right - left, 0.0)
            for (left, right), character in zip(spans, chars)
            if character.strip()
        )
        if not widths:
            return None
        width = widths[len(widths) // 2]
        if width <= 0:
            return None

        gaps = [spans[index + 1][0] - spans[index][1] for index in range(len(spans) - 1)]
        ordered = sorted(gaps)

        # Kalau baris ini sudah punya spasi yang dikenali, lebar kotaknya adalah
        # ukuran spasi yang sebenarnya pada tulisan itu, jadi dipakai sebagai
        # patokan. Selain itu ambangnya sengaja tinggi: kotak per karakter cukup
        # berisik, dan spasi yang telanjur disisipkan di tengah kata lebih
        # merepotkan daripada spasi yang tetap hilang.
        space_widths = sorted(
            right - left
            for (left, right), character in zip(spans, chars)
            if not character.strip()
        )
        if space_widths:
            base = space_widths[len(space_widths) // 2] * 0.75
        else:
            base = width * 0.72
        threshold = max(ordered[len(ordered) // 2] + width * 0.5, base)

        parts = [chars[0]]
        for index, gap in enumerate(gaps):
            neighbours_are_glyphs = chars[index].strip() and chars[index + 1].strip()
            if gap > threshold and neighbours_are_glyphs:
                parts.append(" ")
            parts.append(chars[index + 1])

        respaced = "".join(parts)
        return respaced if respaced != "".join(chars) else None

    @staticmethod
    def _require_ocr_reader():
        """Ambil pembaca OCR, atau tolak permintaan dengan pesan yang jelas."""
        reader = ConvertService._load_ocr_reader()
        if reader is None:
            # ValueError agar endpoint membalas 400: ini kondisi dokumen,
            # bukan kegagalan server
            raise ValueError(
                "PDF ini hasil pindaian dan tidak punya lapisan teks, jadi isinya "
                "harus dikenali lewat OCR. Fitur OCR belum aktif di server ini."
            )
        return reader

    @staticmethod
    def _guard_ocr_page_count(page_count: int) -> None:
        """Tolak dokumen pindaian yang terlalu panjang untuk di-OCR."""
        if page_count > OCR_MAX_PAGES:
            raise ValueError(
                f"PDF ini hasil pindaian dengan {page_count} halaman, terlalu panjang "
                f"untuk dikenali lewat OCR (batas {OCR_MAX_PAGES} halaman). "
                "Pecah dokumennya lebih dulu."
            )

    @staticmethod
    def _ocr_timeout(page_count: int) -> float:
        """Batas waktu OCR dihitung per halaman karena jauh lebih lambat."""
        return max(PROCESS_TIMEOUT, page_count * OCR_TIMEOUT_PER_PAGE)

    @staticmethod
    def _ocr_canvas_items(canvas, reader) -> list:
        """
        Kenali isi satu kanvas halaman, kotaknya dalam satuan titik.

        Koordinat piksel dikembalikan ke satuan halaman, sehingga kotaknya bisa
        langsung dipakai untuk menaruh teks di posisi yang sama.
        """
        items = reader(canvas.to_png())

        scale = 1.0 / canvas.scale
        for item in items:
            x0, y0, x1, y1 = item["bbox"]
            item["bbox"] = (x0 * scale, y0 * scale, x1 * scale, y1 * scale)
        return items

    @staticmethod
    def _ocr_page_items(page, reader) -> list:
        """
        Kenali isi satu halaman apa adanya, tanpa pelurusan citra.

        Dipakai fitur yang menempelkan hasilnya kembali ke halaman PDF asli,
        jadi koordinatnya harus tetap berimpit dengan halaman itu.
        """
        return ConvertService._ocr_canvas_items(PageCanvas(page), reader)

    @staticmethod
    def _page_image(page):
        """Render satu halaman jadi citra BGR, atau None tanpa OpenCV/NumPy."""
        try:
            import cv2
            import numpy as np
        except ImportError:
            return None

        pixmap = page.get_pixmap(dpi=OCR_DPI)
        buffer = np.frombuffer(pixmap.samples, dtype=np.uint8)
        image = buffer.reshape(pixmap.height, pixmap.stride)[:, : pixmap.width * pixmap.n]
        image = image.reshape(pixmap.height, pixmap.width, pixmap.n)

        if pixmap.n == 4:
            return cv2.cvtColor(image, cv2.COLOR_RGBA2BGR)
        if pixmap.n == 3:
            return cv2.cvtColor(image, cv2.COLOR_RGB2BGR)
        return cv2.cvtColor(image[:, :, 0], cv2.COLOR_GRAY2BGR)

    @staticmethod
    def _warp_document_quad(image):
        """
        Cari bidang kertas di dalam foto lalu luruskan perspektifnya.

        Foto dokumen hampir selalu menyertakan meja atau tangan di sekelilingnya,
        dan sisi kertasnya menjorok karena kamera tidak tegak lurus. Keduanya
        membuat baris teks melengkung dan latarnya ikut terbaca sebagai gambar.
        Kembalikan None kalau bidang kertasnya tidak ketemu meyakinkan, supaya
        pindaian datar tidak ikut diubah-ubah.
        """
        import cv2
        import numpy as np

        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
        height, width = gray.shape
        ratio = PHOTO_EDGE_WORK_SIZE / max(height, width)
        if ratio >= 1.0:
            ratio = 1.0
        small = cv2.resize(gray, None, fx=ratio, fy=ratio, interpolation=cv2.INTER_AREA)

        edges = cv2.Canny(cv2.GaussianBlur(small, (5, 5), 0), 50, 150)
        edges = cv2.dilate(edges, np.ones((3, 3), np.uint8))
        contours, _ = cv2.findContours(edges, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

        total_area = small.shape[0] * small.shape[1]
        quad = None
        for contour in sorted(contours, key=cv2.contourArea, reverse=True)[:6]:
            area = cv2.contourArea(contour)
            if not PHOTO_QUAD_MIN_AREA <= area / total_area <= PHOTO_QUAD_MAX_AREA:
                continue
            approx = cv2.approxPolyDP(contour, 0.02 * cv2.arcLength(contour, True), True)
            if len(approx) != 4 or not cv2.isContourConvex(approx):
                continue
            quad = approx.reshape(4, 2).astype(np.float32) / ratio
            break

        if quad is None:
            return None

        # Urutkan sudut: kiri atas, kanan atas, kanan bawah, kiri bawah
        total = quad.sum(axis=1)
        diagonal = np.diff(quad, axis=1).ravel()
        corners = np.array(
            [
                quad[np.argmin(total)],
                quad[np.argmin(diagonal)],
                quad[np.argmax(total)],
                quad[np.argmax(diagonal)],
            ],
            dtype=np.float32,
        )

        def side(first, second) -> float:
            return float(np.linalg.norm(corners[first] - corners[second]))

        target_width = max(side(0, 1), side(3, 2))
        target_height = max(side(0, 3), side(1, 2))
        if target_width < 200 or target_height < 200:
            return None

        target = np.array(
            [
                [0, 0],
                [target_width - 1, 0],
                [target_width - 1, target_height - 1],
                [0, target_height - 1],
            ],
            dtype=np.float32,
        )
        matrix = cv2.getPerspectiveTransform(corners, target)
        return cv2.warpPerspective(
            image,
            matrix,
            (int(round(target_width)), int(round(target_height))),
            flags=cv2.INTER_CUBIC,
            borderMode=cv2.BORDER_REPLICATE,
        )

    @staticmethod
    def _page_skew_angle(image) -> float:
        """
        Perkirakan kemiringan halaman dari arah baris teksnya.

        Kata-kata disatukan dulu jadi gumpalan sepanjang baris, lalu sudut tiap
        gumpalan panjang diambil dan dicari nilai tengahnya. Cara ini tahan
        terhadap logo dan tanda tangan yang arahnya semrawut.
        """
        import cv2

        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
        ratio = PHOTO_SKEW_WORK_SIZE / max(gray.shape)
        if ratio < 1.0:
            gray = cv2.resize(gray, None, fx=ratio, fy=ratio, interpolation=cv2.INTER_AREA)

        binary = cv2.threshold(
            gray, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU
        )[1]
        blobs = cv2.dilate(binary, cv2.getStructuringElement(cv2.MORPH_RECT, (25, 3)))
        contours, _ = cv2.findContours(blobs, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

        angles = []
        for contour in contours:
            (_, _), (box_width, box_height), angle = cv2.minAreaRect(contour)
            if box_width < box_height:
                box_width, box_height = box_height, box_width
                angle += 90.0
            angle = (angle + 45.0) % 90.0 - 45.0
            if box_width < 40 or box_height < 2 or box_width / max(box_height, 1) < 5:
                continue
            if abs(angle) > DESKEW_MAX_ANGLE:
                continue
            angles.append(angle)

        if len(angles) < 5:
            return 0.0

        angles.sort()
        return angles[len(angles) // 2]

    @staticmethod
    def _rotate_image(image, angle: float):
        """Putar citra sebesar sudut tertentu tanpa memotong sudut-sudutnya."""
        import cv2

        height, width = image.shape[:2]
        matrix = cv2.getRotationMatrix2D((width / 2.0, height / 2.0), angle, 1.0)
        cosine, sine = abs(matrix[0, 0]), abs(matrix[0, 1])
        rotated_width = int(round(height * sine + width * cosine))
        rotated_height = int(round(height * cosine + width * sine))
        matrix[0, 2] += rotated_width / 2.0 - width / 2.0
        matrix[1, 2] += rotated_height / 2.0 - height / 2.0
        return cv2.warpAffine(
            image,
            matrix,
            (rotated_width, rotated_height),
            flags=cv2.INTER_CUBIC,
            borderMode=cv2.BORDER_REPLICATE,
        )

    @staticmethod
    def _flatten_illumination(image):
        """
        Ratakan cahaya halaman supaya kertasnya kembali putih.

        Foto selalu punya sisi yang lebih gelap karena bayangan tangan atau
        badan. Tanpa diratakan, ambang tinta menganggap seluruh sisi gelap itu
        sebagai gambar, dan hasilnya DOCX penuh potongan bayangan. Kembalikan
        None kalau cahayanya memang sudah rata.
        """
        import cv2
        import numpy as np

        gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
        height, width = gray.shape
        small = cv2.resize(gray, (max(width // 12, 8), max(height // 12, 8)),
                           interpolation=cv2.INTER_AREA)
        background = cv2.medianBlur(small, 21)
        spread = int(background.max()) - int(background.min())
        if spread < PHOTO_SHADOW_SPREAD:
            return None

        background = cv2.resize(background, (width, height), interpolation=cv2.INTER_CUBIC)
        background = np.maximum(background, 1).astype(np.float32)
        # Pembagian dilakukan per kanal warna supaya stempel dan logo berwarna
        # tidak ikut luntur jadi abu-abu
        flattened = image.astype(np.float32) * (255.0 / background[:, :, None])
        return np.clip(flattened, 0, 255).astype(np.uint8)

    @staticmethod
    def _points_per_pixel(image) -> float:
        """
        Tentukan berapa titik halaman yang diwakili satu piksel bidang kertas.

        Dipakai setelah perspektif diluruskan: kotak halaman PDF-nya seukuran
        bidang foto, bukan seukuran kertasnya, jadi skala aslinya tidak lagi
        berlaku. Kertas dianggap selebar A4 supaya ukuran huruf di Word masuk
        akal, sedangkan tingginya mengikuti perbandingan citra apa adanya.
        """
        height, width = image.shape[:2]
        if height >= width:
            return 595.0 / max(width, 1)
        return 842.0 / max(width, 1)

    @staticmethod
    def _build_page_canvas(page, normalize: bool) -> "PageCanvas":
        """
        Siapkan kanvas halaman, sekalian dibereskan kalau berasal dari foto.

        Citranya selalu dibuat karena deteksi logo dan tanda tangan memerlukannya.
        Pelurusan geometrinya yang dibatasi: hanya untuk halaman hasil foto, sebab
        PDF pindaian yang sudah punya lapisan teks memakai koordinat halaman
        aslinya, dan citra yang diputar tidak lagi berimpit dengan koordinat itu.
        """
        image = ConvertService._page_image(page)
        if image is None:
            return PageCanvas(page)
        if not normalize:
            return PageCanvas(page, image)

        # Rotasi hanya memperbesar bidang citra, tidak mengubah ukuran fisik
        # kertasnya, jadi skala awalnya tetap dipakai selama perspektifnya tidak
        # ikut diluruskan.
        points_per_pixel = 72.0 / OCR_DPI
        try:
            warped = ConvertService._warp_document_quad(image)
            if warped is not None:
                image = warped
                points_per_pixel = ConvertService._points_per_pixel(image)
                logger.info("Bidang kertas pada foto ditemukan dan diluruskan")

            angle = ConvertService._page_skew_angle(image)
            if abs(angle) >= DESKEW_MIN_ANGLE:
                height, width = image.shape[:2]
                rotated = ConvertService._rotate_image(image, angle)
                # Memutar citra memperbesar bidangnya, padahal kertasnya tidak
                # ikut membesar; bidangnya dikembalikan ke ukuran semula supaya
                # halaman Word-nya tidak jadi lebih besar daripada kertas asli
                top = max((rotated.shape[0] - height) // 2, 0)
                left = max((rotated.shape[1] - width) // 2, 0)
                image = rotated[top:top + height, left:left + width]
                logger.info(f"Halaman dimiringkan balik {angle:.1f} derajat")

            flattened = ConvertService._flatten_illumination(image)
            if flattened is not None:
                image = flattened
                logger.info("Cahaya halaman diratakan")
        except Exception:
            logger.warning("Pembersihan citra halaman gagal, dipakai apa adanya", exc_info=True)
            return PageCanvas(page)

        height, width = image.shape[:2]
        return PageCanvas(
            page, image, width * points_per_pixel, height * points_per_pixel
        )

    @staticmethod
    def _group_items_into_lines(items: list) -> list:
        """
        Kelompokkan potongan teks menjadi baris sesuai urutan baca.

        Penggabungan memakai tumpang tindih vertikal, bukan jarak titik tengah:
        pada kop surat ukuran huruf antarpotongan bisa jauh berbeda, sehingga
        patokan titik tengah memecah satu baris menjadi beberapa potongan lepas.
        """
        boxed = [item for item in items if item.get("text", "").strip()]
        if not boxed:
            return []

        lines: list = []
        for item in sorted(boxed, key=lambda entry: (entry["bbox"][1], entry["bbox"][0])):
            _, y0, _, y1 = item["bbox"]
            height = max(y1 - y0, 0.01)
            for line in lines:
                line_height = max(line["bottom"] - line["top"], 0.01)
                overlap = min(y1, line["bottom"]) - max(y0, line["top"])
                shortest = max(min(height, line_height), 0.01)
                ratio = height / line_height
                # Kotak yang jauh lebih tinggi (logo yang terbaca sebagai satu
                # huruf, misalnya) tidak boleh menarik baris-baris di sebelahnya
                # menjadi satu paragraf
                if overlap >= shortest * 0.45 and 0.45 <= ratio <= 2.2:
                    line["items"].append(item)
                    line["top"] = min(line["top"], y0)
                    line["bottom"] = max(line["bottom"], y1)
                    break
            else:
                lines.append({"items": [item], "top": y0, "bottom": y1})

        for line in lines:
            line["items"].sort(key=lambda entry: entry["bbox"][0])
            line["left"] = min(entry["bbox"][0] for entry in line["items"])
            line["right"] = max(entry["bbox"][2] for entry in line["items"])
            heights = sorted(
                entry["bbox"][3] - entry["bbox"][1] for entry in line["items"]
            )
            line["height"] = heights[len(heights) // 2] or 1.0

        lines.sort(key=lambda line: (line["top"], line["left"]))
        return lines

    @staticmethod
    def _drop_graphic_like_items(items: list) -> list:
        """
        Buang potongan OCR yang sebenarnya gambar, bukan tulisan.

        Logo bundar dan stempel kerap dikenali sebagai satu huruf di dalam kotak
        setinggi beberapa baris. Kalau dibiarkan, isinya jadi huruf nyasar di
        tengah kop surat sekaligus menghapus gambarnya dari pencarian gambar
        halaman, karena area teks memang sengaja dikosongkan di sana.
        """
        heights = sorted(item["bbox"][3] - item["bbox"][1] for item in items)
        if len(heights) < 5:
            return items

        median = heights[len(heights) // 2]
        if median <= 0:
            return items

        kept = []
        for item in items:
            height = item["bbox"][3] - item["bbox"][1]
            if height > median * 2.5 and len(item["text"].strip()) <= 2:
                continue
            kept.append(item)
        return kept

    @staticmethod
    def _line_segments(line: dict):
        """
        Hasilkan pasangan (pemisah, potongan) untuk satu baris.

        Jarak antarpotongan yang lebar diterjemahkan jadi tab supaya susunan
        berkolom seperti nomor dan tanggal surat tidak menempel jadi satu.
        """
        previous_right = None
        previous_text = ""
        for item in line["items"]:
            text = item.get("text", "")
            if not text.strip():
                continue

            separator = ""
            if previous_right is not None:
                gap = item["bbox"][0] - previous_right
                # Span PDF bisa terpotong di tengah kata, jadi jarak nol berarti
                # masih satu kata; kotak OCR selalu utuh per kata atau frasa.
                tight = item.get("size") is not None
                if gap > line["height"] * 1.2:
                    separator = "\t"
                elif gap > line["height"] * 0.12 or not tight:
                    separator = " "
                if separator == " " and (
                    previous_text.endswith(" ") or text.startswith(" ")
                ):
                    separator = ""

            yield separator, item
            previous_right = item["bbox"][2]
            previous_text = text

    @staticmethod
    def _line_text(line: dict) -> str:
        """Rangkai satu baris jadi teks biasa."""
        parts = []
        for separator, item in ConvertService._line_segments(line):
            parts.append(separator)
            parts.append(item["text"])
        return "".join(parts).strip()

    @staticmethod
    def _normalize_ocr_line_heights(lines: list) -> None:
        """
        Seragamkan tinggi baris hasil OCR supaya ukuran hurufnya tidak lompat-lompat.

        Kotak OCR mengikuti bentuk tulisannya: baris tanpa huruf turun seperti
        "j" atau "p" terbaca lebih pendek walau ukurannya sama. Tanpa penyetaraan
        ini satu paragraf bisa berisi tiga ukuran huruf berbeda dan terlihat
        berantakan, padahal aslinya seragam.
        """
        measured = []
        for line in lines:
            if any(item.get("size") for item in line["items"]):
                continue
            text = ConvertService._line_text(line)
            # Baris kapital semua tidak punya huruf turun, jadi kotaknya lebih
            # pendek daripada ukuran huruf sebenarnya
            if text and not any(character.islower() for character in text):
                line["height"] *= 1.22
            measured.append(line)

        if len(measured) < 3:
            return

        heights = sorted(line["height"] for line in measured)
        median = heights[len(heights) // 2]
        if median <= 0:
            return

        for line in measured:
            if 0.85 <= line["height"] / median <= 1.18:
                line["height"] = median

    @staticmethod
    def _font_size_pt(item: dict, line_height: float) -> float:
        """
        Tentukan ukuran huruf sebuah potongan teks.

        Span PDF membawa ukuran aslinya; hasil OCR tidak, jadi ukurannya
        diperkirakan dari tinggi kotak pengenalan.
        """
        size = item.get("size") or line_height * OCR_FONT_HEIGHT_RATIO
        return max(6.0, min(36.0, round(float(size) * 2) / 2))

    @staticmethod
    def _detect_graphic_regions(canvas, text_boxes: list) -> list:
        """
        Cari bagian halaman yang berisi gambar, bukan teks.

        Pada PDF hasil pindaian, logo kop surat, tanda tangan, dan stempel
        menyatu dengan citra halaman sehingga tidak bisa diambil sebagai objek
        gambar. Karena itu halaman dirender, area yang sudah dikenali sebagai
        teks dihapus, lalu sisa tintanya dikelompokkan jadi kotak-kotak gambar.
        """
        try:
            import cv2
            import numpy as np
        except ImportError:
            logger.debug("OpenCV/NumPy tidak tersedia, gambar halaman dilewati")
            return []

        image = canvas.gray()
        if image is None:
            return []

        scale = canvas.scale
        shrink = GRAPHIC_DPI / (scale * 72.0)
        if shrink < 1.0:
            image = cv2.resize(image, None, fx=shrink, fy=shrink, interpolation=cv2.INTER_AREA)
            scale *= shrink

        height, width = image.shape
        # Ambang tinta dihitung dari sebaran terang halaman itu sendiri: pada
        # foto, kertas jarang benar-benar putih sehingga ambang tetap membuat
        # separuh halaman terbaca sebagai tinta.
        otsu, _ = cv2.threshold(image, 0, 255, cv2.THRESH_BINARY | cv2.THRESH_OTSU)
        threshold_value = max(90, min(int(otsu), GRAPHIC_INK_THRESHOLD))
        ink = (image < threshold_value).astype(np.uint8)

        pad = max(1, int(round(2 * scale)))
        for x0, y0, x1, y1 in text_boxes:
            left = max(int(x0 * scale) - pad, 0)
            top = max(int(y0 * scale) - pad, 0)
            right = min(int(x1 * scale) + pad + 1, width)
            bottom = min(int(y1 * scale) + pad + 1, height)
            if right > left and bottom > top:
                ink[top:bottom, left:right] = 0

        # Bingkai gelap di tepi kertas adalah bayangan pemindai, bukan isi surat
        border = max(1, int(round(6 * scale)))
        ink[:border, :] = 0
        ink[height - border:, :] = 0
        ink[:, :border] = 0
        ink[:, width - border:] = 0

        # Goresan tanda tangan terputus-putus; ditutup dulu supaya jadi satu objek
        kernel_size = max(3, int(round(5 * scale)))
        kernel = np.ones((kernel_size, kernel_size), np.uint8)
        merged = cv2.morphologyEx(ink, cv2.MORPH_CLOSE, kernel)

        count, _, stats, _ = cv2.connectedComponentsWithStats(merged, 8)
        page_area = max(width * height, 1)
        minimum_pixels = GRAPHIC_MIN_AREA_PT * scale * scale

        regions = []
        for index in range(1, count):
            x, y, box_width, box_height, area = (
                int(value) for value in stats[index][:5]
            )
            if area < minimum_pixels:
                continue

            width_pt = box_width / scale
            height_pt = box_height / scale
            # Garis pemisah kop surat panjang tapi tipis, jadi diloloskan terpisah
            rule_line = width_pt >= 48 and height_pt >= 1.5
            if (width_pt < 12 or height_pt < 12) and not rule_line:
                continue
            box_area = max(box_width * box_height, 1)
            covers_page = box_width >= width * 0.85 and box_height >= height * 0.85
            if covers_page or box_area / page_area > 0.8:
                # Sekujur halaman: citra pindaiannya sendiri, bukan satu gambar
                continue
            if area / box_area < 0.02:
                continue

            regions.append(
                (
                    x / scale,
                    y / scale,
                    (x + box_width) / scale,
                    (y + box_height) / scale,
                )
            )

        regions.sort(key=lambda region: (region[1], region[0]))
        return regions[:GRAPHIC_MAX_REGIONS]

    @staticmethod
    def _page_margins(canvas, lines: list, graphics: list) -> tuple:
        """
        Perkirakan margin halaman dari sebaran isinya.

        Margin ini yang jadi acuan rata kiri/tengah/kanan dan indentasi, jadi
        posisi tiap baris di Word mengikuti posisinya di dokumen asli.
        """
        lefts = [line["left"] for line in lines] + [region[0] for region in graphics]
        rights = [line["right"] for line in lines] + [region[2] for region in graphics]
        tops = [line["top"] for line in lines] + [region[1] for region in graphics]
        bottoms = [line["bottom"] for line in lines] + [region[3] for region in graphics]

        def clamp(value: float) -> float:
            return max(18.0, min(108.0, float(value)))

        if not lefts:
            return 72.0, 72.0, 72.0, 72.0

        return (
            clamp(min(lefts)),
            clamp(canvas.width_pt - max(rights)),
            clamp(min(tops)),
            clamp(canvas.height_pt - max(bottoms)),
        )

    @staticmethod
    def _apply_block_position(
        paragraph, left: float, right: float, margins: tuple, page_width: float
    ) -> None:
        """Tentukan rata teks dan indentasi sebuah blok dari posisinya di halaman."""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        margin_left, margin_right = margins[0], margins[1]
        content_right = page_width - margin_right
        content_width = max(content_right - margin_left, 1.0)
        tolerance = max(12.0, content_width * 0.03)
        indent = left - margin_left
        block_center = (left + right) / 2
        content_center = (margin_left + content_right) / 2

        if abs(block_center - content_center) <= tolerance and indent > 24:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif right >= content_right - tolerance and indent > 72:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if indent > 4:
                paragraph.paragraph_format.left_indent = Pt(
                    min(indent, max(content_width - 40, 0))
                )

    @staticmethod
    def _append_layout_page(
        docx_document, canvas, lines: list, graphics: list, margins: tuple
    ) -> None:
        """Tulis satu halaman ke DOCX dengan urutan, posisi, dan gambarnya."""
        from docx.shared import Pt

        blocks = [("text", line["top"], line["bottom"], line) for line in lines]
        blocks += [("image", region[1], region[3], region) for region in graphics]
        blocks.sort(
            key=lambda block: (
                block[1],
                block[3][0] if block[0] == "image" else block[3]["left"],
            )
        )

        page_width = canvas.width_pt
        content_width = max(page_width - margins[0] - margins[1], 1.0)
        previous_bottom = None

        for kind, top, bottom, payload in blocks:
            paragraph = docx_document.add_paragraph()
            spacing = paragraph.paragraph_format
            spacing.space_after = Pt(0)
            spacing.line_spacing = 1.0
            gap = 0.0 if previous_bottom is None else top - previous_bottom
            spacing.space_before = Pt(max(0.0, min(gap, 72.0)))

            if kind == "image":
                image_bytes = canvas.crop_png(payload)
                if image_bytes is None:
                    continue
                left, _, right, _ = payload
                ConvertService._apply_block_position(
                    paragraph, left, right, margins, page_width
                )
                run = paragraph.add_run()
                try:
                    run.add_picture(
                        io.BytesIO(image_bytes),
                        width=Pt(min(right - left + 4.0, content_width)),
                    )
                except Exception:
                    logger.debug("Gambar halaman gagal disisipkan", exc_info=True)
            else:
                ConvertService._apply_block_position(
                    paragraph, payload["left"], payload["right"], margins, page_width
                )
                for separator, item in ConvertService._line_segments(payload):
                    if separator:
                        paragraph.add_run(separator)
                    run = paragraph.add_run(item["text"])
                    run.font.size = Pt(
                        ConvertService._font_size_pt(item, payload["height"])
                    )
                    flags = item.get("flags")
                    if flags is not None:
                        run.bold = bool(flags & 2 ** 4)
                        run.italic = bool(flags & 2 ** 1)

            previous_bottom = bottom

    @staticmethod
    def _build_layout_docx(
        input_path: str,
        output_path: str,
        item_source: Callable,
        progress: Optional[Callable] = None,
        normalize: bool = False,
    ) -> None:
        """
        Susun DOCX yang meniru tata letak PDF, lengkap dengan gambarnya.

        Dipakai untuk PDF hasil pindaian: teksnya bisa datang dari OCR atau dari
        lapisan teks tak terlihat, sedangkan logo kop surat, tanda tangan, dan
        stempel diambil sebagai potongan gambar halaman. Tanpa ini keluarannya
        cuma tumpukan paragraf rata kiri tanpa satu gambar pun.
        """
        fitz = ConvertService._import_fitz()
        from docx import Document
        from docx.shared import Pt

        document = fitz.open(input_path)
        docx_document = Document()
        margins = None
        try:
            total_pages = document.page_count
            for page_index, page in enumerate(document):
                canvas = ConvertService._build_page_canvas(page, normalize)
                items = ConvertService._drop_graphic_like_items(item_source(canvas))
                lines = ConvertService._group_items_into_lines(items)
                ConvertService._normalize_ocr_line_heights(lines)
                graphics = ConvertService._detect_graphic_regions(
                    canvas, [item["bbox"] for item in items]
                )

                if margins is None:
                    margins = ConvertService._page_margins(canvas, lines, graphics)
                    section = docx_document.sections[0]
                    section.page_width = Pt(canvas.width_pt)
                    section.page_height = Pt(canvas.height_pt)
                    section.left_margin = Pt(margins[0])
                    section.right_margin = Pt(margins[1])
                    section.top_margin = Pt(margins[2])
                    section.bottom_margin = Pt(margins[3])
                else:
                    docx_document.add_page_break()

                ConvertService._append_layout_page(
                    docx_document, canvas, lines, graphics, margins
                )
                ConvertService._report_progress(progress, page_index + 1, total_pages)

            docx_document.save(output_path)
        finally:
            document.close()

    @staticmethod
    def _pdf_text_items(canvas) -> list:
        """Ambil span teks halaman lengkap dengan kotak, ukuran, dan gayanya."""
        items = []
        for block in canvas.page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:  # 0 = blok teks
                continue
            for line in block.get("lines") or ():
                for span in line.get("spans") or ():
                    text = span.get("text", "")
                    bbox = span.get("bbox")
                    if not text.strip() or not bbox:
                        continue
                    items.append(
                        {
                            "text": text,
                            "bbox": tuple(bbox),
                            "size": span.get("size"),
                            "flags": span.get("flags", 0),
                        }
                    )
        return items

    @staticmethod
    def _ocr_page_lines(page, reader) -> list:
        """
        Kenali satu halaman lalu kembalikan barisnya sebagai teks biasa.

        Halamannya diluruskan lebih dulu karena keluarannya hanya teks: pada
        foto yang miring, baris kiri dan kanan berbeda tinggi sehingga tanpa
        pelurusan potongan dari dua baris bisa tercampur jadi satu kalimat.
        """
        canvas = ConvertService._build_page_canvas(page, True)
        lines = ConvertService._group_items_into_lines(
            ConvertService._drop_graphic_like_items(
                ConvertService._ocr_canvas_items(canvas, reader)
            )
        )
        return [
            text for text in (ConvertService._line_text(line) for line in lines) if text
        ]

    @staticmethod
    def _ocr_pdf_to_docx(
        input_path: str, output_path: str, progress: Optional[Callable] = None
    ) -> None:
        """
        Kenali teks dari PDF hasil pindaian lalu tulis ke DOCX.

        Dipakai saat PDF sama sekali tidak punya lapisan teks, jadi satu-satunya
        cara menghasilkan Word yang bisa diedit adalah membaca gambar halamannya.
        """
        reader = ConvertService._require_ocr_reader()
        ConvertService._build_layout_docx(
            input_path,
            output_path,
            lambda canvas: ConvertService._ocr_canvas_items(canvas, reader),
            progress,
            normalize=True,
        )

    @staticmethod
    def _report_progress(progress: Optional[Callable], done: int, total: int) -> None:
        """
        Laporkan kemajuan per halaman tanpa membiarkan errornya menggagalkan konversi.

        Callback datang dari pemanggil (mis. job store) dan hanya bersifat
        informatif, jadi kegagalannya tidak boleh membatalkan pekerjaan asli.
        """
        if progress is None:
            return
        try:
            progress(done, total)
        except Exception:
            logger.debug("Callback progres gagal", exc_info=True)

    @staticmethod
    async def pdf_to_docx(
        input_path: str, output_path: str, progress: Optional[Callable] = None
    ) -> bool:
        """
        Konversi PDF ke DOCX dengan isi yang tetap bisa diedit.

        pdf2docx dipakai lebih dulu karena layoutnya jauh lebih terjaga daripada
        LibreOffice, tapi ia hanya membaca satu jenis lapisan teks sekaligus.
        Karena itu modenya dipilih dari hasil analisis PDF, keluarannya diperiksa
        ulang, dan kalau teksnya tetap kosong konversi dijatuhkan ke pembangun
        DOCX berbasis teks (atau OCR untuk PDF hasil pindaian) supaya hasilnya
        tidak pernah berupa gambar saja.
        """
        try:
            import pdf2docx  # noqa: F401
        except ImportError:
            logger.error("pdf2docx tidak terpasang")
            raise RuntimeError("Fitur PDF ke Word belum tersedia di server ini")

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        # Analisis dilakukan di luar konversi karena hasilnya ikut menentukan
        # batas waktu: jalur OCR butuh jauh lebih lama daripada pdf2docx.
        try:
            stats = await asyncio.to_thread(ConvertService._analyze_pdf_text, input_path)
        except Exception as e:
            logger.warning(f"Analisis lapisan teks PDF gagal, pakai mode standar: {e}")
            stats = {"pages": 1, "chars": 0, "chars_per_page": 999.0, "hidden_ratio": 0.0}

        needs_ocr = stats["chars_per_page"] < PDF_MIN_CHARS_PER_PAGE
        if needs_ocr:
            ConvertService._guard_ocr_page_count(stats["pages"])

        timeout = (
            ConvertService._ocr_timeout(stats["pages"]) if needs_ocr else PROCESS_TIMEOUT
        )

        def perform_conversion():
            # Tanpa lapisan teks sama sekali hanya OCR yang bisa menolong
            if needs_ocr:
                logger.info(
                    f"PDF terdeteksi hasil pindaian ({stats['pages']} halaman), "
                    "konversi lewat OCR"
                )
                ConvertService._ocr_pdf_to_docx(input_path, output_path, progress)
                return

            minimum_chars = max(1, int(stats["chars"] * PDF_DOCX_TEXT_RATIO))

            # Teks tersembunyi mendominasi berarti PDF pindaian yang sudah di-OCR.
            # pdf2docx cuma bisa membaca satu lapisan: dengan ocr=2 teksnya dapat
            # tapi kop surat, tanda tangan, dan stempel hilang; dengan ocr=0 yang
            # tersisa cuma gambar halaman. Karena itu dokumen semacam ini disusun
            # sendiri dari span teks PDF plus potongan gambar halamannya.
            if stats["hidden_ratio"] >= PDF_HIDDEN_TEXT_RATIO:
                logger.info(
                    "PDF pindaian dengan lapisan teks OCR, disusun ulang beserta "
                    "gambar halamannya"
                )
                try:
                    ConvertService._build_layout_docx(
                        input_path,
                        output_path,
                        ConvertService._pdf_text_items,
                        progress,
                    )
                except Exception as e:
                    logger.warning(f"Penyusunan ulang tata letak gagal: {e}")
                else:
                    extracted = ConvertService._docx_text_length(output_path)
                    if extracted < 0 or extracted >= minimum_chars:
                        return
                    logger.warning(
                        f"DOCX tata letak cuma berisi {extracted} karakter "
                        f"(PDF punya {stats['chars']}), coba pdf2docx"
                    )

            primary_mode = 2 if stats["hidden_ratio"] >= PDF_HIDDEN_TEXT_RATIO else 0
            modes = [primary_mode, 0 if primary_mode == 2 else 2]

            for mode in modes:
                try:
                    ConvertService._run_pdf2docx(input_path, output_path, mode)
                except Exception as e:
                    logger.warning(f"pdf2docx gagal pada mode ocr={mode}: {e}")
                    continue

                extracted = ConvertService._docx_text_length(output_path)
                if extracted < 0 or extracted >= minimum_chars:
                    return

                logger.warning(
                    f"DOCX dari pdf2docx mode ocr={mode} cuma berisi {extracted} karakter "
                    f"(PDF punya {stats['chars']}), coba mode lain"
                )

            # pdf2docx tetap tidak menghasilkan teks memadai: susun ulang dari span
            # teks PDF supaya isinya dijamin bisa diedit
            logger.info("Membangun DOCX langsung dari teks PDF sebagai cadangan")
            ConvertService._build_docx_from_pdf_text(input_path, output_path)

        try:
            await asyncio.wait_for(
                asyncio.to_thread(perform_conversion), timeout=timeout
            )
        except asyncio.TimeoutError:
            logger.error(f"PDF ke DOCX timeout setelah {timeout:.0f}s")
            return False
        except (ValueError, RuntimeError):
            raise
        except Exception as e:
            logger.error(f"Error saat konversi PDF ke DOCX: {e}", exc_info=True)
            return False

        return os.path.exists(output_path)

    @staticmethod
    async def pdf_to_xlsx(
        input_path: str, output_path: str, progress: Optional[Callable] = None
    ) -> bool:
        """
        Ekstrak tabel dari PDF ke XLSX (satu sheet per tabel).

        LibreOffice tidak mendukung arah ini, jadi dipakai deteksi tabel PyMuPDF.
        Bila tak ada tabel, isinya dijatuhkan ke teks per baris; untuk PDF hasil
        pindaian teks itu diambil lewat OCR supaya keluarannya tidak berupa buku
        kerja kosong yang tetap dilaporkan sukses.
        """
        fitz = ConvertService._import_fitz()
        try:
            from openpyxl import Workbook
        except ImportError:
            logger.error("openpyxl tidak terpasang")
            raise RuntimeError("Fitur PDF ke Excel belum tersedia di server ini")

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        stats = await asyncio.to_thread(ConvertService._analyze_pdf_text, input_path)
        needs_ocr = stats["chars_per_page"] < PDF_MIN_CHARS_PER_PAGE
        if needs_ocr:
            ConvertService._guard_ocr_page_count(stats["pages"])
        timeout = (
            ConvertService._ocr_timeout(stats["pages"]) if needs_ocr else PROCESS_TIMEOUT
        )

        def perform_conversion():
            document = fitz.open(input_path)
            workbook = Workbook()
            workbook.remove(workbook.active)
            table_count = 0

            try:
                total_pages = document.page_count
                for page_index, page in enumerate(document, start=1):
                    try:
                        tables = page.find_tables()
                    except Exception:
                        continue

                    for table_index, table in enumerate(tables.tables, start=1):
                        rows = table.extract()
                        if not rows:
                            continue

                        table_count += 1
                        title = (
                            f"Hal {page_index}"
                            if table_index == 1
                            else f"Hal {page_index}-{table_index}"
                        )
                        sheet = workbook.create_sheet(title[:31])
                        for row in rows:
                            sheet.append(
                                ["" if cell is None else str(cell) for cell in row]
                            )

                    ConvertService._report_progress(progress, page_index, total_pages)

                # Tidak ada tabel terdeteksi: jatuhkan ke teks per baris agar hasil tetap berguna
                if table_count == 0:
                    reader = ConvertService._require_ocr_reader() if needs_ocr else None
                    sheet = workbook.create_sheet("Teks")
                    for page_index, page in enumerate(document, start=1):
                        lines = page.get_text("text").splitlines()
                        if not any(line.strip() for line in lines) and reader is not None:
                            lines = ConvertService._ocr_page_lines(page, reader)

                        for line in lines:
                            sheet.append([f"Hal {page_index}", line])

                        ConvertService._report_progress(progress, page_index, total_pages)

                    if sheet.max_row <= 1 and sheet["A1"].value is None:
                        # ValueError agar endpoint membalas 400: ini kondisi dokumen,
                        # bukan kegagalan server
                        raise ValueError(
                            "Tidak ada tabel maupun teks yang bisa diambil dari PDF ini. "
                            "Kalau dokumennya hasil pindaian, aktifkan OCR di server."
                        )

                workbook.save(output_path)
            finally:
                document.close()

        try:
            await asyncio.wait_for(
                asyncio.to_thread(perform_conversion), timeout=timeout
            )
        except asyncio.TimeoutError:
            logger.error(f"PDF ke XLSX timeout setelah {timeout:.0f}s")
            return False
        except (ValueError, RuntimeError):
            raise
        except Exception as e:
            logger.error(f"Error saat konversi PDF ke XLSX: {e}", exc_info=True)
            return False

        return os.path.exists(output_path)

    @staticmethod
    async def pdf_to_pptx(input_path: str, output_path: str, dpi: int = 150) -> bool:
        """
        Konversi PDF ke PPTX.

        Jalur utama: impor PDF lewat LibreOffice Draw (hasil masih bisa diedit).
        Fallback: render tiap halaman jadi gambar full-slide memakai python-pptx.
        """
        output_dir = os.path.dirname(output_path) or "."
        os.makedirs(output_dir, exist_ok=True)

        produced = await ConvertService._run_libreoffice(
            input_path,
            output_dir,
            [
                "--infilter=draw_pdf_import",
                "--convert-to",
                "pptx:Impress MS PowerPoint 2007 XML",
                "--outdir",
                output_dir,
                input_path,
            ],
            ".pptx",
            "PDF to PPTX",
        )

        if produced:
            if os.path.abspath(produced) != os.path.abspath(output_path):
                shutil.move(produced, output_path)
            return True

        logger.warning(
            "Impor PDF via LibreOffice Draw gagal, memakai fallback berbasis gambar"
        )
        return await ConvertService._pdf_to_pptx_images(input_path, output_path, dpi)

    @staticmethod
    async def _pdf_to_pptx_images(input_path: str, output_path: str, dpi: int) -> bool:
        fitz = ConvertService._import_fitz()
        try:
            from pptx import Presentation
            from pptx.util import Emu
        except ImportError:
            logger.error("python-pptx tidak terpasang")
            raise RuntimeError("Fitur PDF ke PowerPoint belum tersedia di server ini")

        def perform_conversion():
            document = fitz.open(input_path)
            presentation = Presentation()
            blank_layout = presentation.slide_layouts[6]
            zoom = dpi / 72.0

            try:
                with tempfile.TemporaryDirectory() as workdir:
                    for page_index, page in enumerate(document):
                        rect = page.rect
                        # Samakan ukuran slide dengan halaman pertama (1 pt = 12700 EMU)
                        if page_index == 0:
                            presentation.slide_width = Emu(int(rect.width * 12700))
                            presentation.slide_height = Emu(int(rect.height * 12700))

                        pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
                        image_path = os.path.join(workdir, f"page_{page_index}.png")
                        pixmap.save(image_path)

                        slide = presentation.slides.add_slide(blank_layout)
                        slide.shapes.add_picture(
                            image_path,
                            0,
                            0,
                            width=presentation.slide_width,
                            height=presentation.slide_height,
                        )

                    presentation.save(output_path)
            finally:
                document.close()

        try:
            await asyncio.wait_for(
                asyncio.to_thread(perform_conversion), timeout=PROCESS_TIMEOUT
            )
        except asyncio.TimeoutError:
            logger.error("PDF ke PPTX timeout")
            return False
        except Exception as e:
            logger.error(f"Error saat konversi PDF ke PPTX: {e}", exc_info=True)
            return False

        return os.path.exists(output_path)

    # ------------------------------------------------------------------
    # PDF -> gambar / teks
    # ------------------------------------------------------------------
    @staticmethod
    async def pdf_to_images(
        input_path: str,
        output_dir: str,
        image_format: str = "png",
        dpi: int = 150,
        pages: Optional[str] = None,
        base_name: str = "page",
    ) -> list[str]:
        """Render halaman PDF menjadi PNG/JPG. Mengembalikan daftar path gambar."""
        fitz = ConvertService._import_fitz()
        os.makedirs(output_dir, exist_ok=True)

        extension = resolve_image_extension(image_format)

        def perform_conversion() -> list[str]:
            document = fitz.open(input_path)
            try:
                selected = parse_page_ranges(pages, document.page_count)
                zoom = max(dpi, 36) / 72.0
                matrix = fitz.Matrix(zoom, zoom)
                results: list[str] = []

                for page_index in selected:
                    page = document.load_page(page_index)
                    pixmap = page.get_pixmap(matrix=matrix)
                    image_path = os.path.join(
                        output_dir, f"{base_name}_{page_index + 1:03d}.{extension}"
                    )
                    ConvertService._save_pixmap(pixmap, image_path, extension, 90)
                    results.append(image_path)

                return results
            finally:
                document.close()

        return await asyncio.wait_for(
            asyncio.to_thread(perform_conversion), timeout=PROCESS_TIMEOUT
        )

    @staticmethod
    async def pdf_extract_images(
        input_path: str,
        output_dir: str,
        image_format: str = "jpg",
        pages: Optional[str] = None,
        base_name: str = "gambar",
    ) -> list[str]:
        """
        Ambil gambar yang tertanam di dalam PDF apa adanya.

        Berbeda dengan pdf_to_images yang merender seluruh halaman, fungsi ini
        hanya mengeluarkan objek gambar asli - berguna untuk memanen foto dari
        dokumen tanpa ikut membawa teks dan latar halaman.
        """
        fitz = ConvertService._import_fitz()
        os.makedirs(output_dir, exist_ok=True)

        extension = resolve_image_extension(image_format)

        def perform_extraction() -> list[str]:
            document = fitz.open(input_path)
            try:
                selected = parse_page_ranges(pages, document.page_count)
                results: list[str] = []
                # Satu gambar bisa dipakai di banyak halaman; xref menjaga agar
                # tidak tersimpan berulang kali
                seen: set[int] = set()

                for page_index in selected:
                    page = document.load_page(page_index)
                    for image in page.get_images(full=True):
                        xref = image[0]
                        if xref in seen:
                            continue
                        seen.add(xref)

                        pixmap = fitz.Pixmap(document, xref)
                        try:
                            # CMYK dan gambar bertopeng alfa perlu diubah ke RGB
                            # sebelum bisa ditulis sebagai JPG/PNG biasa
                            if pixmap.n - pixmap.alpha >= 4 or (
                                extension == "jpg" and pixmap.alpha
                            ):
                                converted = fitz.Pixmap(fitz.csRGB, pixmap)
                                pixmap = converted

                            image_path = os.path.join(
                                output_dir,
                                f"{base_name}_{page_index + 1:03d}_{len(results) + 1:02d}.{extension}",
                            )
                            ConvertService._save_pixmap(
                                pixmap, image_path, extension, 95
                            )
                            results.append(image_path)
                        finally:
                            pixmap = None

                return results
            finally:
                document.close()

        return await asyncio.wait_for(
            asyncio.to_thread(perform_extraction), timeout=PROCESS_TIMEOUT
        )

    @staticmethod
    async def pdf_to_text(
        input_path: str,
        output_path: str,
        text_format: str = "txt",
        pages: Optional[str] = None,
        progress: Optional[Callable] = None,
    ) -> bool:
        """
        Ekstrak isi PDF menjadi .txt atau .md.

        PDF hasil pindaian tidak punya teks yang bisa diambil, jadi halaman yang
        kosong dikenali lewat OCR. Tanpa itu keluarannya berupa berkas nol byte
        yang tetap dilaporkan sukses.
        """
        fitz = ConvertService._import_fitz()
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        stats = await asyncio.to_thread(ConvertService._analyze_pdf_text, input_path)
        needs_ocr = stats["chars_per_page"] < PDF_MIN_CHARS_PER_PAGE
        if needs_ocr:
            ConvertService._guard_ocr_page_count(stats["pages"])
        timeout = (
            ConvertService._ocr_timeout(stats["pages"]) if needs_ocr else PROCESS_TIMEOUT
        )

        def perform_conversion():
            # pymupdf4llm menghasilkan Markdown yang jauh lebih rapi bila tersedia,
            # tapi ia sama-sama bergantung pada lapisan teks sehingga dilewati untuk
            # dokumen pindaian
            if text_format == "md" and not needs_ocr:
                try:
                    import pymupdf4llm

                    content = pymupdf4llm.to_markdown(input_path)
                    Path(output_path).write_text(content, encoding="utf-8")
                    return
                except ImportError:
                    logger.info(
                        "pymupdf4llm tidak tersedia, memakai ekstraksi Markdown sederhana"
                    )

            reader = ConvertService._require_ocr_reader() if needs_ocr else None

            document = fitz.open(input_path)
            try:
                selected = parse_page_ranges(pages, document.page_count)
                blocks: list[str] = []

                for position, page_index in enumerate(selected, start=1):
                    page = document.load_page(page_index)
                    text = page.get_text("text").strip()
                    if not text and reader is not None:
                        lines = ConvertService._ocr_page_lines(page, reader)
                        text = "\n".join(lines).strip()

                    if text_format == "md":
                        blocks.append(f"## Halaman {page_index + 1}\n\n{text}")
                    else:
                        blocks.append(text)

                    ConvertService._report_progress(progress, position, len(selected))

                if not any(block.strip() for block in blocks):
                    raise ValueError(
                        "Tidak ada teks yang bisa diambil dari PDF ini. Kalau "
                        "dokumennya hasil pindaian, aktifkan OCR di server."
                    )

                separator = "\n\n---\n\n" if text_format == "md" else "\n\n\f\n\n"
                Path(output_path).write_text(separator.join(blocks), encoding="utf-8")
            finally:
                document.close()

        try:
            await asyncio.wait_for(
                asyncio.to_thread(perform_conversion), timeout=timeout
            )
        except asyncio.TimeoutError:
            logger.error(f"Ekstraksi teks PDF timeout setelah {timeout:.0f}s")
            return False
        except (ValueError, RuntimeError):
            raise
        except Exception as e:
            logger.error(f"Error saat ekstraksi teks PDF: {e}", exc_info=True)
            return False

        return os.path.exists(output_path)

    @staticmethod
    async def pdf_to_html(
        input_path: str, output_path: str, pages: Optional[str] = None
    ) -> bool:
        """
        Ekstrak PDF menjadi satu berkas HTML yang mempertahankan tata letak.

        PyMuPDF mengeluarkan dokumen HTML lengkap per halaman, jadi isi <body>
        tiap halaman diambil lalu dijahit menjadi satu berkas agar hasilnya
        tetap satu dokumen yang bisa dibuka langsung di browser.
        """
        fitz = ConvertService._import_fitz()
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        def perform_conversion():
            document = fitz.open(input_path)
            try:
                selected = parse_page_ranges(pages, document.page_count)
                sections: list[str] = []

                for page_index in selected:
                    page = document.load_page(page_index)
                    body = ConvertService._extract_html_body(page.get_text("html"))
                    sections.append(
                        f'<section class="halaman" id="halaman-{page_index + 1}">\n'
                        f"{body}\n</section>"
                    )

                title = Path(input_path).stem
                content = "\n".join(sections)
                Path(output_path).write_text(
                    ConvertService._wrap_html(title, content, page_break_class="halaman"),
                    encoding="utf-8",
                )
            finally:
                document.close()

        try:
            await asyncio.wait_for(
                asyncio.to_thread(perform_conversion), timeout=PROCESS_TIMEOUT
            )
        except asyncio.TimeoutError:
            logger.error("PDF ke HTML timeout")
            return False
        except Exception as e:
            logger.error(f"Error saat konversi PDF ke HTML: {e}", exc_info=True)
            return False

        return os.path.exists(output_path)

    @staticmethod
    async def pdf_to_csv(
        input_path: str,
        output_dir: str,
        pages: Optional[str] = None,
        base_name: str = "tabel",
    ) -> list[str]:
        """
        Ekstrak tiap tabel di PDF menjadi satu berkas CSV.

        Berbeda dengan pdf_to_xlsx yang menumpuk semua tabel dalam satu buku
        kerja, di sini tiap tabel berdiri sendiri supaya mudah diimpor ke alat
        lain. Mengembalikan daftar path CSV yang dihasilkan.
        """
        fitz = ConvertService._import_fitz()
        os.makedirs(output_dir, exist_ok=True)

        def perform_conversion() -> list[str]:
            document = fitz.open(input_path)
            try:
                selected = parse_page_ranges(pages, document.page_count)
                results: list[str] = []

                for page_index in selected:
                    page = document.load_page(page_index)
                    try:
                        tables = page.find_tables()
                    except Exception:
                        continue

                    for table_index, table in enumerate(tables.tables, start=1):
                        rows = table.extract()
                        if not rows:
                            continue

                        csv_path = os.path.join(
                            output_dir,
                            f"{base_name}_{page_index + 1:03d}_{table_index:02d}.csv",
                        )
                        # utf-8-sig supaya Excel di Windows membaca karakter
                        # non-ASCII dengan benar
                        with open(
                            csv_path, "w", newline="", encoding="utf-8-sig"
                        ) as handle:
                            writer = csv.writer(handle)
                            for row in rows:
                                writer.writerow(
                                    ["" if cell is None else str(cell) for cell in row]
                                )
                        results.append(csv_path)

                return results
            finally:
                document.close()

        return await asyncio.wait_for(
            asyncio.to_thread(perform_conversion), timeout=PROCESS_TIMEOUT
        )

    # ------------------------------------------------------------------
    # HTML / Markdown -> PDF
    # ------------------------------------------------------------------
    @staticmethod
    async def html_to_pdf(
        output_path: str,
        html_path: Optional[str] = None,
        url: Optional[str] = None,
        base_url: Optional[str] = None,
    ) -> bool:
        """Render berkas HTML atau URL menjadi PDF memakai WeasyPrint."""
        try:
            from weasyprint import HTML
        except ImportError:
            logger.error("weasyprint tidak terpasang")
            raise RuntimeError("Fitur HTML ke PDF belum tersedia di server ini")
        except OSError as e:
            # WeasyPrint memuat Pango/Cairo lewat cffi saat diimpor. Di mesin tanpa
            # pustaka GTK (mis. Windows dev) impor gagal dengan OSError, bukan ImportError.
            logger.error(f"Pustaka sistem WeasyPrint tidak tersedia: {e}")
            raise RuntimeError(
                "Fitur HTML ke PDF butuh pustaka Pango/Cairo yang belum terpasang di server ini"
            )

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        def perform_conversion():
            if url:
                document = HTML(url=url)
            else:
                document = HTML(filename=html_path, base_url=base_url)
            document.write_pdf(output_path)

        try:
            await asyncio.wait_for(
                asyncio.to_thread(perform_conversion), timeout=PROCESS_TIMEOUT
            )
        except asyncio.TimeoutError:
            logger.error("HTML ke PDF timeout")
            return False
        except Exception as e:
            logger.error(f"Error saat konversi HTML ke PDF: {e}", exc_info=True)
            return False

        return os.path.exists(output_path)

    @staticmethod
    async def markdown_to_pdf(input_path: str, output_path: str) -> bool:
        """Markdown -> HTML -> PDF."""
        try:
            import markdown as markdown_lib
        except ImportError:
            logger.error("markdown tidak terpasang")
            raise RuntimeError("Fitur Markdown ke PDF belum tersedia di server ini")

        source = Path(input_path).read_text(encoding="utf-8", errors="replace")
        body = markdown_lib.markdown(
            source, extensions=["extra", "sane_lists", "tables", "toc"]
        )
        html = ConvertService._wrap_html(Path(input_path).stem, body)

        with tempfile.NamedTemporaryFile(
            "w", suffix=".html", delete=False, encoding="utf-8"
        ) as handle:
            handle.write(html)
            temp_html = handle.name

        try:
            return await ConvertService.html_to_pdf(
                output_path, html_path=temp_html, base_url=str(Path(input_path).parent)
            )
        finally:
            try:
                os.remove(temp_html)
            except OSError:
                pass

    # ------------------------------------------------------------------
    # EPUB
    # ------------------------------------------------------------------
    @staticmethod
    async def epub_to_pdf(input_path: str, output_path: str) -> bool:
        """Gabungkan seluruh dokumen HTML di dalam EPUB lalu render ke PDF."""
        try:
            import ebooklib
            from ebooklib import epub
        except ImportError:
            logger.error("ebooklib tidak terpasang")
            raise RuntimeError("Fitur EPUB ke PDF belum tersedia di server ini")

        def extract_html() -> str:
            book = epub.read_epub(input_path)
            chunks: list[str] = []
            for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                content = item.get_content().decode("utf-8", errors="replace")
                chunks.append(f'<div class="chapter">{content}</div>')
            return "".join(chunks)

        try:
            body = await asyncio.to_thread(extract_html)
        except Exception as e:
            logger.error(f"Gagal membaca EPUB: {e}", exc_info=True)
            return False

        html = ConvertService._wrap_html(
            Path(input_path).stem, body, page_break_class="chapter"
        )

        with tempfile.NamedTemporaryFile(
            "w", suffix=".html", delete=False, encoding="utf-8"
        ) as handle:
            handle.write(html)
            temp_html = handle.name

        try:
            return await ConvertService.html_to_pdf(output_path, html_path=temp_html)
        finally:
            try:
                os.remove(temp_html)
            except OSError:
                pass

    @staticmethod
    async def pdf_to_epub(
        input_path: str,
        output_path: str,
        title: Optional[str] = None,
        progress: Optional[Callable] = None,
    ) -> bool:
        """
        Bangun EPUB dari teks PDF (satu bab per halaman).

        Layout kompleks tidak dipertahankan - EPUB memang format teks mengalir.
        Halaman tanpa teks dikenali lewat OCR; tanpa itu PDF hasil pindaian
        menghasilkan EPUB berisi bab-bab kosong yang tetap dilaporkan sukses.
        """
        fitz = ConvertService._import_fitz()
        try:
            from ebooklib import epub
        except ImportError:
            logger.error("ebooklib tidak terpasang")
            raise RuntimeError("Fitur PDF ke EPUB belum tersedia di server ini")

        import html as html_module

        stats = await asyncio.to_thread(ConvertService._analyze_pdf_text, input_path)
        needs_ocr = stats["chars_per_page"] < PDF_MIN_CHARS_PER_PAGE
        if needs_ocr:
            ConvertService._guard_ocr_page_count(stats["pages"])
        timeout = (
            ConvertService._ocr_timeout(stats["pages"]) if needs_ocr else PROCESS_TIMEOUT
        )

        def perform_conversion():
            reader = ConvertService._require_ocr_reader() if needs_ocr else None
            document = fitz.open(input_path)
            try:
                book = epub.EpubBook()
                book.set_identifier(uuid.uuid4().hex)
                book.set_title(title or Path(input_path).stem)
                book.set_language("id")

                chapters = []
                total_pages = document.page_count
                has_text = False
                for page_index, page in enumerate(document, start=1):
                    text = page.get_text("text").strip()
                    if not text and reader is not None:
                        text = "\n\n".join(
                            ConvertService._ocr_page_lines(page, reader)
                        ).strip()
                    has_text = has_text or bool(text)
                    paragraphs = (
                        "".join(
                            f"<p>{html_module.escape(block)}</p>"
                            for block in text.split("\n\n")
                            if block.strip()
                        )
                        or "<p></p>"
                    )

                    chapter = epub.EpubHtml(
                        title=f"Halaman {page_index}",
                        file_name=f"page_{page_index}.xhtml",
                        lang="id",
                    )
                    chapter.content = f"<h2>Halaman {page_index}</h2>{paragraphs}"
                    book.add_item(chapter)
                    chapters.append(chapter)
                    ConvertService._report_progress(progress, page_index, total_pages)

                if not has_text:
                    # ValueError agar endpoint membalas 400: ini kondisi dokumen,
                    # bukan kegagalan server
                    raise ValueError(
                        "Tidak ada teks yang bisa diambil dari PDF ini, jadi EPUB-nya "
                        "akan kosong. Kalau dokumennya hasil pindaian, aktifkan OCR "
                        "di server."
                    )

                book.toc = tuple(chapters)
                book.add_item(epub.EpubNcx())
                book.add_item(epub.EpubNav())
                book.spine = ["nav", *chapters]

                epub.write_epub(output_path, book)
            finally:
                document.close()

        try:
            await asyncio.wait_for(
                asyncio.to_thread(perform_conversion), timeout=timeout
            )
        except asyncio.TimeoutError:
            logger.error(f"PDF ke EPUB timeout setelah {timeout:.0f}s")
            return False
        except (ValueError, RuntimeError):
            raise
        except Exception as e:
            logger.error(f"Error saat konversi PDF ke EPUB: {e}", exc_info=True)
            return False

        return os.path.exists(output_path)

    # ------------------------------------------------------------------
    # PDF/A
    # ------------------------------------------------------------------
    @staticmethod
    async def pdf_to_pdfa(input_path: str, output_path: str, version: int = 2) -> bool:
        """Konversi PDF biasa ke PDF/A (arsip) memakai Ghostscript."""
        if version not in (1, 2, 3):
            raise ValueError("Versi PDF/A harus 1, 2, atau 3")

        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

        command = [
            "gs",
            f"-dPDFA={version}",
            "-dBATCH",
            "-dNOPAUSE",
            "-dQUIET",
            "-dSAFER",
            "-dPDFACompatibilityPolicy=1",
            "-sColorConversionStrategy=UseDeviceIndependentColor",
            "-sDEVICE=pdfwrite",
            f"-sOutputFile={output_path}",
            input_path,
        ]

        success = await PDFService._execute_command(
            command, f"PDF/A-{version} Conversion"
        )
        return success and os.path.exists(output_path)

    # ------------------------------------------------------------------
    # Gambar
    # ------------------------------------------------------------------
    @staticmethod
    async def normalize_image(input_path: str, output_dir: str) -> str:
        """
        Pastikan gambar bisa dipakai img2pdf.

        HEIC/HEIF dan gambar ber-alpha ditolak img2pdf, jadi dinormalkan dulu
        ke JPEG. Mengembalikan path yang siap dikonversi.
        """
        ext = Path(input_path).suffix.lower()

        if ext in (".jpg", ".jpeg"):
            return input_path

        try:
            from PIL import Image
        except ImportError:
            logger.warning("Pillow tidak terpasang, gambar dipakai apa adanya")
            return input_path

        if ext in (".heic", ".heif"):
            try:
                import pillow_heif

                pillow_heif.register_heif_opener()
            except ImportError:
                logger.error("pillow-heif tidak terpasang")
                raise RuntimeError("Format HEIC belum didukung di server ini")

        os.makedirs(output_dir, exist_ok=True)
        target_path = os.path.join(
            output_dir, f"{Path(input_path).stem}_normalized.jpg"
        )

        def perform_conversion():
            with Image.open(input_path) as image:
                # img2pdf menolak alpha channel; komposit ke latar putih
                if image.mode in ("RGBA", "LA", "P"):
                    rgba = image.convert("RGBA")
                    background = Image.new("RGB", rgba.size, (255, 255, 255))
                    background.paste(rgba, mask=rgba.split()[-1])
                    background.save(target_path, "JPEG", quality=95)
                else:
                    image.convert("RGB").save(target_path, "JPEG", quality=95)

        try:
            await asyncio.to_thread(perform_conversion)
        except Exception as e:
            logger.error(f"Gagal menormalkan gambar {input_path}: {e}")
            return input_path

        return target_path if os.path.exists(target_path) else input_path

    # ------------------------------------------------------------------
    # Helper
    # ------------------------------------------------------------------
    @staticmethod
    def _save_pixmap(
        pixmap, image_path: str, extension: str, jpg_quality: int = 90
    ) -> None:
        """
        Tulis pixmap PyMuPDF ke berkas gambar.

        PyMuPDF hanya bisa menulis JPG dan PNG; WebP dan TIFF dibangun ulang
        lewat Pillow dari buffer piksel mentah agar tidak perlu berkas antara.
        """
        pillow_format = PILLOW_IMAGE_FORMATS.get(extension)

        if pillow_format is None:
            if extension == "jpg":
                pixmap.save(image_path, jpg_quality=jpg_quality)
            else:
                pixmap.save(image_path)
            return

        from PIL import Image

        mode = "RGBA" if pixmap.alpha else "RGB"
        image = Image.frombytes(mode, (pixmap.width, pixmap.height), pixmap.samples)

        if pillow_format == "WEBP":
            image.save(image_path, "WEBP", quality=jpg_quality, method=4)
        else:
            # LZW/deflate menjaga TIFF tetap lossless tanpa membengkak
            image.save(image_path, "TIFF", compression="tiff_deflate")

    @staticmethod
    def _import_fitz():
        # Nama modul "fitz" sudah deprecated; "pymupdf" dipakai lebih dulu
        try:
            import pymupdf

            return pymupdf
        except ImportError:
            pass

        try:
            import fitz

            return fitz
        except ImportError:
            logger.error("PyMuPDF tidak terpasang")
            raise RuntimeError(
                "Fitur ini membutuhkan PyMuPDF yang belum terpasang di server"
            )

    @staticmethod
    async def _run_libreoffice(
        input_path: str,
        output_dir: str,
        command_tail: list[str],
        expected_suffix: str,
        task_name: str,
    ) -> Optional[str]:
        """Jalankan LibreOffice memakai profil dari pool, kembalikan path hasil."""
        async with libreoffice_pool.acquire() as profile_dir:
            command = [
                "libreoffice",
                f"-env:UserInstallation=file://{profile_dir}",
                "--headless",
                "--norestore",
                "--invisible",
                *command_tail,
            ]

            success = await PDFService._execute_command(command, task_name)

            if not success:
                await libreoffice_pool.reset(profile_dir)
                return None

        expected_path = os.path.join(
            output_dir, f"{Path(input_path).stem}{expected_suffix}"
        )
        if os.path.exists(expected_path):
            logger.info(f"{task_name} berhasil: {expected_path}")
            return expected_path

        logger.error(f"{task_name}: output tidak ditemukan di {expected_path}")
        return None

    # Cocokkan isi <body> beserta atributnya, lintas baris
    _BODY_PATTERN = re.compile(
        r"<body[^>]*>(.*)</body>", re.IGNORECASE | re.DOTALL
    )

    @staticmethod
    def _extract_html_body(html: str) -> str:
        """
        Ambil isi <body> dari dokumen HTML.

        Keluaran PyMuPDF adalah dokumen HTML utuh per halaman; tanpa langkah ini
        berkas gabungan akan berisi banyak <html> bersarang.
        """
        match = ConvertService._BODY_PATTERN.search(html)
        return match.group(1).strip() if match else html.strip()

    @staticmethod
    def _wrap_html(
        title: str, body: str, page_break_class: Optional[str] = None
    ) -> str:
        page_break_rule = (
            f".{page_break_class} {{ page-break-after: always; }}"
            if page_break_class
            else ""
        )
        return f"""<!DOCTYPE html>
<html lang="id">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  @page {{ size: A4; margin: 2cm; }}
  body {{ font-family: "Liberation Serif", Georgia, serif; font-size: 11pt; line-height: 1.6; }}
  h1, h2, h3, h4 {{ font-family: "Liberation Sans", Arial, sans-serif; line-height: 1.3; }}
  code, pre {{ font-family: "Liberation Mono", monospace; font-size: 9.5pt; }}
  pre {{ background: #f5f5f5; padding: 10px; border-radius: 4px; white-space: pre-wrap; }}
  table {{ border-collapse: collapse; width: 100%; }}
  th, td {{ border: 1px solid #ddd; padding: 6px; text-align: left; }}
  img {{ max-width: 100%; }}
  {page_break_rule}
</style>
</head>
<body>
{body}
</body>
</html>"""
